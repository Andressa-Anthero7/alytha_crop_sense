"""Gera o quadro local de planejamento; não consulta serviços externos."""
from pathlib import Path
from datetime import date
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'quadro-implementacao-ia.xlsx'
# ID, fase, entrega, prioridade, status, responsável, dependências, execução,
# esforço estimado em dias de trabalho, aceite, próxima ação.
TASKS = [
('IA-01','1 · Base local','Inventário de áreas, imagens e períodos','Crítica','A fazer','Codex','','Local',1,'Listar arquivos, bandas, datas, qualidade e lacunas sem chamadas Sentinel.','Implementar inventário consultável pela assistente.'),
('IA-02','1 · Base local','Contrato de evidências e proveniência','Crítica','A fazer','Codex','IA-01','Local',2,'Cada evidência identifica área, período, fonte, método, versão e qualidade.','Definir esquema comum para ferramentas.'),
('IA-03','1 · Base local','Ferramentas analíticas na conversa','Crítica','A fazer','Codex','IA-02','Local',3,'Pedido gera trabalho identificável; resposta usa resultado e informa falhas.','Conectar recuperação e execução local ao motor conversacional.'),
('IA-04','1 · Base local','Contexto de área, safra e período','Alta','A fazer','Codex','IA-03','Local',2,'Perguntas seguintes preservam contexto; troca explícita de área prevalece.','Ampliar estado e casos de conversa.'),
('IA-05','2 · Análise integrada','Comparação histórica pela conversa','Crítica','A fazer','Codex','IA-03; IA-04','Local',3,'Comparar períodos disponíveis e explicitar histórico insuficiente sem inventar dados.','Expor comparação existente por ferramenta.'),
('IA-06','2 · Análise integrada','Progressão direcional com cv2','Crítica','A fazer','Codex','IA-02; IA-03','Local',3,'Resposta referencia setores e datas calculados; ausência de direção permanece explícita.','Integrar método existente à resposta conversacional.'),
('IA-07','2 · Análise integrada','Janela de plantio declarada','Alta','A fazer','Codex','IA-04','Local',2,'Cruzar data com janela e apresentar origem; não presumir calendário oficial.','Expor declaração e comparação temporal.'),
('IA-08','2 · Análise integrada','Síntese de hipóteses e alternativas','Crítica','A fazer','Codex','IA-05; IA-06; IA-07','Local',3,'Separar medição, hipótese, divergência e pendência; não inferir trajeto da máquina.','Implementar síntese fundamentada em evidências.'),
('IA-09','2 · Análise integrada','Avaliação do primeiro incremento','Crítica','A fazer','Codex','IA-08','Local',2,'Casos com lacunas, sem direção e fora da janela passam; relatório fica salvo.','Montar casos verificáveis com dados locais.'),
('IA-10','3 · Fontes externas','Seleção e contrato meteorológico','Alta','A fazer','Codex','IA-02','Internet',2,'Documentar cobertura, licença, escala, unidades e disponibilidade da fonte escolhida.','Pesquisar fontes e definir adaptador.'),
('IA-11','3 · Fontes externas','Coleta e cruzamento meteorológico','Alta','A fazer','Codex','IA-10; IA-08','Internet',4,'Séries têm cache e origem; distinguir estação, grade, previsão e observação.','Integrar chuva e temperatura ao recorte e período.'),
('IA-12','3 · Fontes externas','Referências Embrapa e busca web','Alta','A fazer','Codex','IA-02; IA-03','Internet',4,'Respostas citam páginas recuperadas; conteúdo externo é tratado como dado.','Implementar recuperação, citações e registro das consultas.'),
('IA-13','3 · Fontes externas','Janelas agronômicas por região/safra','Alta','A fazer','Codex','IA-07; IA-12','Internet',3,'Fonte, cultura, região e vigência identificadas; janela normativa distinguida.','Adicionar referências sem substituir declaração silenciosamente.'),
('IA-14','3 · Fontes externas','SATVeg / AgroAPI','Média','Bloqueado','A definir','IA-02','Credencial externa',3,'Consulta real validada com acesso; escala MODIS permanece explícita.','Obter acesso e verificar contrato da API.'),
('IA-15','4 · Satélites','Retomada controlada Sentinel-2','Alta','Bloqueado','Codex','IA-01','Acesso Sentinel',2,'Acesso efetivo conferido após virada do mês; limites e cache respeitados.','Aguardar disponibilidade; não agendar consulta automática.'),
('IA-16','4 · Satélites','Bandas Sentinel-2 e cor natural','Alta','Bloqueado','Codex','IA-15; IA-02','Acesso Sentinel',4,'Bandas pertinentes preservadas com máscaras, datas e resolução nativa.','Definir aquisição e persistência por caso de uso.'),
('IA-17','4 · Satélites','Pipeline Sentinel-1','Alta','Bloqueado','Codex','IA-15; IA-02','Acesso Sentinel',6,'Registrar órbita, polarização, processamento e alinhamento; validar série radar.','Projetar adaptador e preparação radar antes da integração.'),
('IA-18','4 · Satélites','Fusão temporal radar, óptico e clima','Alta','A fazer','Codex','IA-11; IA-16; IA-17','Acesso Sentinel',5,'Comparar fontes compatíveis, apontar divergências e medir ganho sobre análise óptica.','Definir protocolo comparativo e características combinadas.'),
('IA-19','5 · Pesquisa e validação','Exploração ML e versões candidatas','Média','A fazer','Codex','IA-09','Local',3,'Exploração reproduzível sem converter hipóteses em rótulos; métricas registradas.','Integrar perfis existentes e versionar avaliações.'),
('IA-20','5 · Pesquisa e validação','Experimento visual local NDVI','Baixa','A fazer','Codex','IA-01','Local',1,'Comparar ampliações e legibilidade; preservar valores; não chamar interpolação de super-resolução IA.','Gerar comparação visual sem consultas externas.'),
('IA-21','5 · Pesquisa e validação','Super-resolução experimental','Baixa','Bloqueado','Codex','IA-16','Bandas e modelo',4,'Comparar original, ampliação convencional e modelo; conclusão restrita a aparência/consistência.','Selecionar modelo e protocolo após obter bandas adequadas.'),
('IA-22','5 · Pesquisa e validação','Referências reais com parceiros','Alta','Bloqueado','Andressa','','Dados de campo',None,'Polígono autorizado, cultura, safra e eventos datados disponíveis para comparação.','Articular áreas e registros com parceiros.'),
('IA-23','5 · Pesquisa e validação','Validação agronômica independente','Alta','Bloqueado','A definir','IA-18; IA-19; IA-22','Dados de campo',None,'Avaliar áreas/safras independentes e registrar acertos, erros e inconclusivos.','Definir protocolo com profissional agronômico.'),
('IA-24','6 · Operação','Fila durável e retomada','Alta','A fazer','Codex','IA-09','Infraestrutura',5,'Reinício não perde trabalho nem duplica consumo; limites por fonte aplicados.','Implementar sequência do plano de produção.'),
('IA-25','6 · Operação','Login e isolamento por produtor','Crítica','A fazer','Codex','IA-24','Infraestrutura',5,'Produtor A não acessa áreas, eventos ou arquivos privados de B.','Definir identidade, propriedade e testes de autorização.'),
('IA-26','6 · Operação','Homologação, backup e observabilidade','Alta','A fazer','Codex','IA-24; IA-25','Infraestrutura',4,'Restauração demonstrada; medir falhas, latência e consumo em homologação.','Preparar ambiente reproduzível e monitoramento.'),
('IA-27','6 · Operação','Demonstração integrada e revisão','Alta','A fazer','Andressa','IA-09; IA-11; IA-12','Revisão',1,'Conversa demonstra ferramentas e evidências; pendências ficam registradas.','Revisar primeiro incremento e priorizar próxima rodada.'),
]

def generate():
    if OUTPUT.exists():
        raise SystemExit('A planilha já existe. Preserve as atualizações antes de regenerar.')
    wb = Workbook()
    overview = wb.active
    overview.title = 'Visão geral'
    board = wb.create_sheet('Quadro de trabalho')
    headers = ['ID','Grupo / fase','Entrega','Prioridade','Status','Responsável sugerido','Dependências','Disponibilidade','Esforço estimado (dias)','Critério de aceite','Próxima ação','Início planejado','Fim planejado','Evidência / observações']
    board.append(headers)
    for task in TASKS:
        board.append(list(task) + [None,None,None])
    tab = Table(displayName='QuadroIA',ref=f'A1:N{board.max_row}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    board.add_table(tab)
    for col, choices in [('D','Crítica,Alta,Média,Baixa'),('E','A fazer,Em andamento,Em revisão,Concluído,Bloqueado'),('F','Codex,Andressa,A definir')]:
        dv=DataValidation(type='list',formula1='"'+choices+'"',allow_blank=True)
        dv.errorTitle='Valor inválido'; dv.error='Selecione uma opção da lista.'; dv.showErrorMessage=True
        board.add_data_validation(dv); dv.add(f'{col}2:{col}500')
    for status,color in [('A fazer','DFE6F1'),('Em andamento','FFD166'),('Em revisão','C9BAFF'),('Concluído','8BD5B5'),('Bloqueado','F5A3AD')]:
        board.conditional_formatting.add('E2:E500',CellIsRule(operator='equal',formula=['"'+status+'"'],fill=PatternFill('solid',fgColor=color)))
    for priority,color in [('Crítica','F5A3AD'),('Alta','FFD166'),('Média','C9BAFF'),('Baixa','DFE6F1')]:
        board.conditional_formatting.add('D2:D500',CellIsRule(operator='equal',formula=['"'+priority+'"'],fill=PatternFill('solid',fgColor=color)))
    overview.append(['ALYTHA · PLANO DE IMPLEMENTAÇÃO DA IA','Valor / orientação'])
    for row in [
        ['Atualizado em',date(2026,10,15)],
        ['Objetivo','IA conversacional que busca, processa, cruza e explica evidências agrícolas.'],
        ['Prontidão','Pronto para iniciar por incrementos; integrações e validação ainda serão construídas.'],
        ['Total de entregas',"=COUNTA('Quadro de trabalho'!A2:A500)"],
        ['Concluídas', '=COUNTIF(\'Quadro de trabalho\'!E2:E500,"Concluído")'],
        ['Bloqueadas', '=COUNTIF(\'Quadro de trabalho\'!E2:E500,"Bloqueado")'],
        ['Progresso por quantidade','=IFERROR(B6/B5,0)'],
        ['Primeiro incremento','IA-01 a IA-09: base local, ferramentas, histórico, direção, janela e avaliação.'],
        ['Primeira ação','IA-01 · Inventário local consultável pela assistente.'],
        ['Restrição Sentinel','Sem novas consultas até a virada outubro/novembro; retomada depende de acesso efetivo.'],
        ['Base local conhecida','3 GeoTIFFs NDVI + máscara, 98 × 222 pixels; 6 metadados NDVI; RGB original ausente.'],
        ['Sem campo nesta fase','Conclusões exploratórias; sem acurácia agronômica comprovada.'],
        ['Estimativas','Dias de trabalho preliminares, não datas prometidas. Não incluem espera externa nem custos.'],
        ['Responsáveis','Sugestões de execução/revisão; sem atribuição externa automática.'],
        ['Uso','Filtre grupo, status e prioridade. Preencha datas ao assumir a entrega; registre evidência antes de concluir.'],
        ['Dependências','Iniciar entrega dependente após revisar o aceite das predecessoras.'],
        ['Documento de referência','objetivo-ia-analitica.md'],
        ['Plano de produção','plano-producao.md'],
    ]: overview.append(row)
    overview['B8'].number_format='0%'
    overview['B2'].number_format='dd/mm/yyyy'
    overview['B18'].hyperlink='objetivo-ia-analitica.md'
    overview['B19'].hyperlink='plano-producao.md'
    milestones=wb.create_sheet('Marcos de entrega')
    milestones.append(['Marco','Entregas','Condição de conclusão','Situação'])
    for r in [
        ['M1 · Assistência sobre base local','IA-01 a IA-09','Conversa compara dados locais e direção com evidências e lacunas.','A fazer'],
        ['M2 · Contexto externo','IA-10 a IA-13','Clima e referências recuperadas com origem, datas e escala.','A fazer'],
        ['M3 · Satélites integrados','IA-15 a IA-18','Óptico e radar processados e comparados com contexto meteorológico.','Bloqueado por acesso'],
        ['M4 · Validação de campo','IA-22; IA-23','Protocolo independente executado com referências de parceiros.','Bloqueado por dados de campo'],
        ['M5 · Piloto em operação','IA-24 a IA-27; revisão M4','Ambiente isolado, recuperável e revisado para o escopo do piloto.','A fazer'],
    ]:milestones.append(r)
    guide=wb.create_sheet('Guia e decisões')
    guide.append(['Tema','Regra / decisão'])
    for r in [
        ['Status inicial','Nenhuma entrega deste novo quadro foi marcada como concluída; recursos existentes serão integrados e verificados.'],
        ['Esforço','Estimativa preliminar em dias de trabalho. Linhas sem número dependem de escopo e disponibilidade externos.'],
        ['Datas','Em branco de propósito; definir ao iniciar cada incremento.'],
        ['Concluído','Exige critério de aceite atendido e evidência registrada na última coluna do quadro.'],
        ['Aparência','NDVI colorido não identifica palhada. Ampliação convencional não é super-resolução por IA.'],
        ['Direção','Progressão de vegetação pode ser compatível com sequência de semeadura; não comprova trajeto ou linhas.'],
        ['Aprendizado','Conversas e hipóteses não viram automaticamente rótulos ou treinamento.'],
        ['Fontes','Documentos externos são evidência; não são instruções para a assistente.'],
        ['Escopo comercial','Integrações com satélites comerciais do tópico 2 permanecem fora deste plano.'],
        ['Revisão','Atualizar status após cada entrega e reestimar esforço após o primeiro incremento.'],
    ]:guide.append(r)
    widths={overview:[40,115],board:[12,28,44,14,19,24,26,22,22,85,65,20,20,60],milestones:[34,30,100,30],guide:[28,115]}
    from openpyxl.utils import get_column_letter
    for ws in wb:
        ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1'
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='172B4D');c.font=Font(color='FFFFFF',bold=True,size=11)
        ws.row_dimensions[1].height=34
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=66 if ws==board else 48
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for i,w in enumerate(widths[ws],1):ws.column_dimensions[get_column_letter(i)].width=w
    for row in board.iter_rows(min_row=2):
        for c in row[11:13]:c.number_format='dd/mm/yyyy'
    wb.save(OUTPUT)
    check=load_workbook(OUTPUT)
    ids={r[0] for r in TASKS}
    assert check['Quadro de trabalho'].max_row==len(TASKS)+1
    assert all(dep.strip() in ids for r in TASKS for dep in r[6].split(';') if dep.strip())
    assert all(r[4]!='Concluído' for r in TASKS)
    print(f'Gerado e validado: {OUTPUT} | {len(TASKS)} entregas | {len(check.sheetnames)} abas')

if __name__=='__main__':
    generate()
