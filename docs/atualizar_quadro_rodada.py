"""Update the existing local board after reading verified execution evidence."""
import hashlib
import json
import shutil
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
from uuid import uuid4
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
BOARD=ROOT/'docs/quadro-implementacao-ia.xlsx'
DAY=datetime(2026,10,15)
TAG='Rodada 15/10/2026:'

UPDATES={
 'IA-02':('Concluído','Revisar contratos ao adicionar novas fontes.','Contrato v1 nas ferramentas históricas, espaciais, janela e clima; projeção das evidências antigas explicita qualidade e versões desconhecidas. app/evidence_contract.py; docs/analises-conversacionais.md.'),
 'IA-03':('Concluído','Usar os recibos de trabalho para acompanhar análises.','Ferramentas executam resultados locais sem provedor; recibos persistentes com ID e estado final, erros explícitos e GET /api/assistant/work/{id}. app/analytical_assistant.py; testes e demonstração real salvos.'),
 'IA-04':('Concluído','Revisar continuidade nas conversas de demonstração.','Contexto opaco por até 24h, seleção atual prevalece; cultura/safra/corte explícitos preservados na interface. Testes de mudança de área, corte e safra passaram. app/conversation_context.py.'),
 'IA-05':('Concluído','Revisar a utilidade da comparação com parceiros.','Comparação sazonal exposta na conversa; cobertura e anos insuficientes ficam inconclusivos. Referência real: 4 observações recentes, seis anos comparáveis. data/outputs/assistant/offline-evaluation.json.'),
 'IA-06':('Concluído','Validar interpretação das frentes com dados de campo.','Direção usa resultado espacial salvo ou recalcula três rasters exatos locais. Informa setores/datas ou ausência de direção. Caso real não detectou direção; teste com frente datada passou. app/local_spatial.py.'),
 'IA-07':('Concluído','Informar janelas reais do produtor quando disponíveis.','Declaração pela conversa com datas e origem, histórico imutável, cultura/safra, comparação dentro/fora e declaração retrospectiva. Não presume calendário oficial. Testes passaram.'),
 'IA-08':('Concluído','Revisar hipóteses com referências independentes.','Síntese separa medições, hipóteses/alternativas, divergências, lacunas e verificação seguinte. Nenhuma trajetória de máquina é inferida; clima em cache pode complementar. Demonstração real salva.'),
 'IA-09':('Concluído','Submeter a demonstração à revisão humana em IA-27.','Avaliação técnica offline salva com regressão completa e quatro casos sobre o histórico real; inclui ausência de direção, lacunas, baixa cobertura e fora da janela. data/outputs/assistant/offline-evaluation.json.'),
 'IA-10':('Concluído','Revisar condições da fonte ao ampliar a operação.','NASA POWER selecionada: reanálise de grade 0,5° × 0,625°, diária UTC, chuva mm/dia e temperatura °C; origem, política de uso/citação, revisões e disponibilidade documentadas em analises-conversacionais.md.'),
 'IA-11':('Concluído','Comparar a grade com medições de campo quando disponíveis.','Adaptador NASA POWER com cache por célula/período, unidades verificadas, sentinelas e lacunas. Consulta real de 30 dias e cruzamento com janelas NDVI; distingue grade de estação/previsão. app/weather.py.'),
 'IA-12':('Concluído','Ampliar fontes somente com contratos de recuperação verificáveis.','Busca web restrita ao portal oficial Embrapa, até três páginas efetivamente recuperadas, links, hashes, cache e registro da consulta. Conteúdo tratado como dado. Não equivale a busca web irrestrita nem revisão do texto integral. app/reference_search.py.'),
 'IA-13':('Em andamento','Carregar referências revisadas e confirmar solo, cultivar, manejo e risco aplicáveis ao piloto.','Importador e comparação regional implementados/testados, com origem, cultura, safra, região, vigência e tipos separados. Portaria oficial MT 2026/27 localizada. Sem carga de intervalos aplicáveis: condições do talhão ainda desconhecidas. app/regional_calendar.py.'),
 'IA-14':('Bloqueado','Renovar/corrigir acesso AgroAPI e repetir validação real.','Token presente no .env, sem exposição de conteúdo. Consulta real em Polygon local compatível retornou HTTP 401; nenhum dado SATVeg foi simulado. data/outputs/assistant/satveg-access.json.'),
 'IA-18':('Bloqueado','Retomar após acesso e séries ópticas/radar compatíveis; definir protocolo comparativo.','Clima disponível em IA-11. Fusão e medição de ganho continuam sem bandas Sentinel-1/2 compatíveis e validação. Respeitada restrição de novas consultas Sentinel até novembro.'),
 'IA-19':('Concluído','Comparar candidatos com outras áreas/safras independentes quando disponíveis.','Experimento real versionado: 226 janelas, 158 iniciais e 68 posteriores, KMeans k=2/3/4, sementes e bibliotecas registradas; silhouette, estabilidade e distâncias. Sem rótulos de manejo nem promoção automática. app/ml_experiment.py; data/outputs/experiments/ml.'),
 'IA-20':('Concluído','Revisar legibilidade das visualizações geradas.','Comparação local nearest/bilinear/bicúbica em RGB; original 98×222, 21.248 pixels válidos. Hash antes/depois idêntico; máscara reaplicada. Sem super-resolução IA. data/outputs/experiments/ndvi; app/ndvi_visual_experiment.py.'),
 'IA-24':('Em andamento','Migrar coletores externos e aplicar limites/deduplicação por fornecedor.','Fila durável local com migration 5, worker separado, deduplicação, reservas renováveis, retomada e proteção contra conclusão obsoleta. Testes de concorrência/reinício e smoke HTTP→worker→resultado passaram. Coletores externos ainda não migrados. docs/operacao-local.md.'),
 'IA-25':('Em andamento','Definir provedor de identidade e revisar organizações, usuários e reivindicação dos registros existentes.','Modelo de identidade/propriedade e testes de autorização necessários documentados. Login, isolamento e autorização em endpoints/arquivos/workers ainda não implementados. Não atribuir dados antigos ao primeiro usuário. docs/operacao-local.md.'),
 'IA-26':('Em andamento','Executar build Docker, verificar volumes, configurar backup externo e medir operação completa.','Backup core real restaurado em pasta nova: 61 arquivos, integridade, hashes e contagens por tabela conferidos. Endpoint de saúde e métricas locais; dependências fixadas e receita Docker/Compose. Docker ausente: ambiente Linux não homologado. docs/operacao-local.md.'),
 'IA-27':('Em revisão','Andressa revisar a demonstração e registrar utilidade, falhas e prioridades.','Quatro conversas reais locais salvas; clima e pesquisa Embrapa validados externamente, navegação/layout verificados no navegador. Resultado técnico pronto para revisão humana; nenhuma revisão de Andressa presumida. data/outputs/assistant/offline-evaluation.json.'),
}


def update():
    report=json.loads((ROOT/'data/outputs/assistant/offline-evaluation.json').read_text(encoding='utf-8'))
    if not report['tests']['passed'] or report['tests']['total']<304:
        raise SystemExit('A avaliação final precisa passar antes de concluir as entregas.')
    original=hashlib.sha256(BOARD.read_bytes()).hexdigest()
    wb=load_workbook(BOARD);ws=wb['Quadro de trabalho']
    rows={ws.cell(i,1).value:i for i in range(2,ws.max_row+1)}
    if not set(UPDATES)<=set(rows):raise SystemExit('IDs do quadro não correspondem à atualização.')
    backup=ROOT/'local_backups'/('quadro-antes-da-rodada-'+uuid4().hex[:12]+'.xlsx')
    backup.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(BOARD,backup)
    for identifier,(status,action,evidence) in UPDATES.items():
        row=rows[identifier]
        ws.cell(row,5,status);ws.cell(row,11,action)
        note=ws.cell(row,14).value or ''
        if TAG not in note:ws.cell(row,14,(note+'\n' if note else '')+TAG+' '+evidence)
        if not ws.cell(row,12).value:ws.cell(row,12,DAY)
        if status=='Concluído' and not ws.cell(row,13).value:ws.cell(row,13,DAY)
        ws.row_dimensions[row].height=100
    counts=Counter(ws.cell(r,5).value for r in rows.values())
    overview=wb['Visão geral'];overview['B2']=DAY
    overview['B4']=f"{counts['Concluído']} de {len(rows)} entregas concluídas; {counts['Em andamento']} em andamento, {counts['Em revisão']} em revisão e {counts['Bloqueado']} bloqueadas. Acurácia agronômica e produção permanecem pendentes."
    overview['B10']='IA-13: referências regionais revisadas; IA-24 a IA-26: operação e isolamento; IA-27: revisão humana.'
    milestones=wb['Marcos de entrega']
    for row,status in ((2,'Concluído tecnicamente; validação agronômica pendente'),(3,'Em andamento · IA-13 parcial'),(4,'Bloqueado por acesso Sentinel'),(5,'Bloqueado por dados de campo'),(6,'Em andamento · operação parcial')):
        milestones.cell(row,4,status)
    guide=wb['Guia e decisões'];guide['A2']='Atualização da rodada'
    guide['B2']='Status revistos em 15/10/2026 com testes, execução real e evidências; itens parciais, dependências externas e revisão humana permanecem explícitos.'
    log=wb['Registro 15-10'] if 'Registro 15-10' in wb.sheetnames else wb.create_sheet('Registro 15-10')
    if log.max_row==1 and log['A1'].value is None:
        log.append(['ID','Status da rodada','Evidência / limite'])
        for c in log[1]:c._style=copy(ws['A1']._style)
        for identifier,(status,_,evidence) in UPDATES.items():log.append([identifier,status,evidence])
        log.freeze_panes='A2';log.sheet_view.showGridLines=False
        for column,width in [('A',14),('B',23),('C',115)]:log.column_dimensions[column].width=width
        for row in log.iter_rows(min_row=2):
            log.row_dimensions[row[0].row].height=75
            for cell in row:cell.alignment=copy(ws['N2'].alignment)
    wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
    temporary=BOARD.with_name('quadro-atualizando-'+uuid4().hex+'.xlsx')
    try:
        wb.save(temporary)
        check=load_workbook(temporary)
        assert check['Quadro de trabalho'].max_row==29
        assert Counter(check['Quadro de trabalho'].cell(r,5).value for r in range(2,30))==counts
        assert counts['Concluído']==14 and counts['Bloqueado']==9
        if hashlib.sha256(BOARD.read_bytes()).hexdigest()!=original:raise SystemExit('A planilha mudou durante a atualização; cópia preparada preservada.')
        temporary.replace(BOARD)
    finally:temporary.unlink(missing_ok=True)
    print(json.dumps({'board':str(BOARD),'backup':str(backup),'status_counts':dict(counts),'tests':report['tests']['total']},ensure_ascii=True))


if __name__=='__main__':update()
