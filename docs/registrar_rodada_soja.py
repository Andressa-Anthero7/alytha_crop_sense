"""Append the authorized soybean research round without regenerating the board."""
import hashlib
import json
import shutil
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
BOARD=ROOT/'docs/quadro-implementacao-ia.xlsx'
TAG='Rodada 16/10/2026 · ciclo da soja:'


def update():
    catalogue=json.loads((ROOT/'web/soy-study.json').read_text(encoding='utf-8'))
    lines=['# Ciclo da soja · matriz inicial de estudo','',
           'Pesquisa inicial em 16/10/2026. Esta matriz descreve o que estudar e implementar; não diagnostica a área selecionada. O catálogo pode ser consultado no painel de camadas, em **Ciclo da soja · estudo e planejamento**.','',
           'As fases fenológicas e os manejos transversais não são uma sequência de dias fixa. Cultivar, ambiente e manejo alteram o ciclo; manejos podem ocorrer em várias fases. O rótulo automático de aumento do verde permanece **sinal — possível plantio ou rebrota**.','',
           '## Temas e evidências','']
    sources={s['id']:s for s in catalogue['sources']}
    for group in dict.fromkeys(t['group'] for t in catalogue['topics']):
        lines += ['### '+group,'','| Tema | Pergunta | Proposta no sistema | Informação de campo / limite |','|---|---|---|---|']
        for topic in catalogue['topics']:
            if topic['group']!=group:continue
            refs=', '.join('['+sources[i]['title']+']('+sources[i]['url']+')' for i in topic['sources'])
            lines.append('| '+topic['name']+' | '+topic['question']+' | '+topic['system']+' | '+topic['field']+' '+topic['limit']+' '+refs+' |')
        lines.append('')
    lines += ['## Próximas entregas','']
    for feature in catalogue['features']:lines+=['- **'+feature['title']+'** — '+feature['status']+'. '+feature['next']]
    lines+=['','Pesquisa, protocolo e critérios de aceite: [rodada-2026-10-16-soja.md](rodada-2026-10-16-soja.md). Nenhum reconhecimento de etapa, regra agronômica ou modelo é considerado validado apenas por constar nesta matriz.','']
    (ROOT/'docs/ciclo-soja-estudo.md').write_text('\n'.join(lines),encoding='utf-8')
    wb=load_workbook(BOARD);sheet=wb['Quadro de trabalho']
    rows={sheet.cell(i,1).value:i for i in range(2,sheet.max_row+1)}
    updates={
        'IA-13':('Priorizar calendário sanitário e ZARC aplicáveis à data pretendida; exigir município, safra, solo, grupo de cultivar, manejo e risco antes de concluir dentro/fora.',
                 'Escopo ampliado para decisão de plantio. Reutilizar app/regional_calendar.py; não tratar declaração do usuário como conferência oficial. docs/rodada-2026-10-16-soja.md.'),
        'IA-17':('Preparar contrato e testes radar; após liberação de aquisição, validar séries VV/VH de órbita compatível, processamento, qualidade e alinhamento com óptico/clima.',
                 'Prioridade: condição da superfície e mudanças compatíveis com umidade, sem conversão universal de dB em umidade ou aptidão. Pesquisa ESA/Copernicus registrada; aquisição permanece bloqueada até acesso autorizado/efetivo.'),
        'IA-18':('Comparar óptico + clima com óptico + clima + radar; medir utilidade para condição de plantio, sequência de cobertura e inconclusivos por etapa.',
                 'Protocolo proposto: mesma área, safra, disponibilidade temporal, qualidade e validação de campo. Radar não define janela oficial nem confirma plantio. Não iniciar coleta automática.'),
        'IA-23':('Definir validação por etapa, com áreas/safras independentes e registros datados; incluir janela, condição de solo, plantio/rebrota e causas alternativas.',
                 'Acurácia agronômica permanece pendente. Catálogo e regras exploratórias não são validação de diagnóstico, produtividade ou manejo.'),
    }
    for identifier,(action,note) in updates.items():
        row=rows[identifier];sheet.cell(row,11,action)
        old=sheet.cell(row,14).value or ''
        if TAG not in old:sheet.cell(row,14,(old+'\n' if old else '')+TAG+' '+note)
    tasks=[
        ('IA-29','3 · Fontes externas','Janela e radar na conclusão de condição para plantio','Alta','A fazer','Codex','IA-13; IA-17; IA-18','Referências oficiais + acesso radar',None,
         'Conclusão separa janela sanitária, ZARC aplicável e condição física/climática; cada fator tem fonte, vigência e qualidade. Ausência não vira aptidão; comparar com observações independentes.',
         'Conectar referências revisadas ao card; em seguida incorporar evidência radar validada, preservando a conclusão simples e as lacunas.',None,None,
         TAG+' Incluído a pedido do usuário. Triagem clima/solo/NDVI já existe; integração oficial e Sentinel-1 ainda não concluídas. docs/condicoes-plantio.md; docs/rodada-2026-10-16-soja.md.'),
        ('IA-30','5 · Pesquisa e validação','Pesquisa e catálogo do ciclo completo da soja','Alta','Em andamento','Codex','IA-12','Local + referências oficiais',None,
         'Catalogar etapas e manejos transversais com fontes, evidências possíveis, limites e dados de campo; revisão agronômica antes de regras operacionais.',
         'Revisar 22 temas do catálogo com responsável agronômico; aprofundar primeiro janela/solo/implantação e preparar casos reais.',None,None,
         TAG+' Pesquisa inicial e catálogo navegável entregues; estudo aprofundado, revisão agronômica e reconhecimento automático por etapa permanecem pendentes. web/soy-study.json; docs/ciclo-soja-estudo.md.'),
        ('IA-31','2 · Análise integrada','Acompanhamento por etapa da safra e manejos','Alta','A fazer','Codex','IA-30; IA-29; IA-22; IA-23','Dados de campo + validação',None,
         'Cada etapa tem entrada, regra versionada, conclusão curta, alternativas, teste positivo/negativo/lacuna e validação independente. Não inferir R1–R8, praga, doença ou aplicação só por NDVI.',
         'Definir contrato de safra e eventos por etapa; implementar primeiro implantação/estande e água, com evidências e ação de verificação.',None,None,
         TAG+' Escopo do planejamento à pós-colheita; não equivale a treinamento concluído nem a identificação automática disponível.'),
    ]
    for task in tasks:
        if task[0] in rows:continue
        template=sheet.max_row;sheet.append(task);row=sheet.max_row
        for col in range(1,15):
            sheet.cell(row,col)._style=copy(sheet.cell(template,col)._style)
            sheet.cell(row,col).alignment=Alignment(vertical='top',wrap_text=True)
        sheet.row_dimensions[row].height=150
    for table in sheet.tables.values():
        table.ref=f'A1:N{sheet.max_row}'
        if table.autoFilter:table.autoFilter.ref=table.ref
    overview=wb.worksheets[0];counts=Counter(sheet.cell(i,5).value for i in range(2,sheet.max_row+1))
    overview.cell(2,2,datetime(2026,10,16))
    overview.cell(4,2,f"{counts['Concluído']} de {sheet.max_row-1} entregas concluídas; {counts['Em andamento']} em andamento, {counts['Em revisão']} em revisão, {counts['Bloqueado']} bloqueadas e {counts['A fazer']} a fazer. Pesquisa do ciclo inicial; validação agronômica pendente.")
    study_name='Ciclo soja 16-10'
    if study_name not in wb.sheetnames:
        study=wb.create_sheet(study_name)
        study.append(['Tema','Grupo','Etapa / manejo','Pergunta','Objetivo','Proposta no sistema','Dados de campo','Limite','Fontes','Tarefas','Situação'])
        sources={s['id']:s for s in catalogue['sources']}
        for topic in catalogue['topics']:
            study.append([topic['id'],topic['group'],topic['name'],topic['question'],topic['objective'],topic['system'],topic['field'],topic['limit'], '\n'.join(sources[i]['url'] for i in topic['sources']),topic['work'],'Pesquisa inicial · não é diagnóstico'])
        study.freeze_panes='D2';study.auto_filter.ref=study.dimensions
        for row in study:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in study[1]:cell.fill=PatternFill('solid',fgColor='24563D');cell.font=Font(color='FFFFFF',bold=True)
        for col,width in {'A':22,'B':25,'C':34,'D':45,'E':35,'F':65,'G':55,'H':65,'I':55,'J':28,'K':35}.items():study.column_dimensions[col].width=width
        for row in range(2,study.max_row+1):study.row_dimensions[row].height=120
    record_name='Rodada 16-10 soja'
    if record_name not in wb.sheetnames:
        record=wb.create_sheet(record_name)
        record.append(['Data','Escopo','Estado / evidência'])
        for feature in catalogue['features']:record.append(['16/10/2026',feature['title'],feature['status']+' · '+feature['next']])
        record.append(['16/10/2026','Limite','Pesquisa inicial; nenhum treinamento ou diagnóstico por etapa concluído. Sem aquisição Sentinel nesta rodada.'])
        record.append(['16/10/2026','Referência','docs/rodada-2026-10-16-soja.md; docs/ciclo-soja-estudo.md'])
        for col in ['A','B','C']:record.column_dimensions[col].width=24 if col=='A' else 55 if col=='B' else 110
        for row in record:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        record.freeze_panes='A2'
    wb.calculation.fullCalcOnLoad=True
    backup=ROOT/'local_backups'/('quadro-antes-soja-16-10-'+uuid4().hex[:12]+'.xlsx')
    backup.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(BOARD,backup)
    temp=BOARD.with_name('quadro-soja-'+uuid4().hex[:12]+'.xlsx');wb.save(temp)
    check=load_workbook(temp)
    assert {'IA-29','IA-30','IA-31'} <= {check['Quadro de trabalho'].cell(i,1).value for i in range(2,check['Quadro de trabalho'].max_row+1)}
    assert check[study_name].max_row==len(catalogue['topics'])+1
    temp.replace(BOARD)
    print(json.dumps({'board':str(BOARD),'backup':str(backup),'sha256':hashlib.sha256(BOARD.read_bytes()).hexdigest(),'tasks':31,'topics':22,'counts':dict(counts)},ensure_ascii=True))


if __name__=='__main__':update()
